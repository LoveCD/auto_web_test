# -*- coding: utf-8 -*-
import os
from docx import Document

d = "docs/case-docs"
f = [x for x in os.listdir(d) if x.startswith("测试报告-INTL-REAL")][-1]
doc = Document(os.path.join(d, f))
print(f"===== {f} =====")
print("段落数:", len(doc.paragraphs), " 表格数:", len(doc.tables), " 内联图片数:", len(doc.inline_shapes))
print("\n===== 段落 =====")
for i, p in enumerate(doc.paragraphs):
    style = p.style.name
    has_img = "graphic" in p._p.xml
    txt = p.text.strip()
    img_mark = " [IMG]" if has_img else ""
    if txt or has_img:
        print(f"[{i}][{style}]{img_mark} {txt[:75]}")
